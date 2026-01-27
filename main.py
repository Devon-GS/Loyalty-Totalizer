import os
from openpyxl import load_workbook
from tkinter import filedialog
from tkinter import messagebox


try:
	# Specify the folder path where the Excel files are located
	folder_path = filedialog.askdirectory(
		title="Select Cash Up Folder", initialdir=os.getcwd()
	)

	# Collect all Excel files in the folder
	excel_files = [f for f in os.listdir(folder_path) if f.endswith(".xlsx")]

	# Check for excel temp files
	excel_files_filtered = []
	for x in excel_files:
		if not x.startswith("~$"):
			excel_files_filtered.append(x)

	# Specify the output file path
	output_file = filedialog.askopenfilename(
		title="Select Loyaty File", initialdir=os.getcwd()
)
except FileNotFoundError:
	messagebox.showerror(title='File Not Found', message='No File Selected or Found. Please Try Again!')
	os.sys.exit()

# Load the output workbook
output_workbook = load_workbook(output_file)

# Starting row for writing data
output_worksheet = output_workbook.active
start_row = 2
start_column = 2  # Column B

# Loop through each Excel file and extract data
for file in excel_files_filtered:
    file_path = os.path.join(folder_path, file)

    # Get Cell F9 Value
    current_workbook = load_workbook(file_path, data_only=True)
    current_worksheet = current_workbook.active
    cell_f9_value = current_worksheet['F9'].value
    current_workbook.save(file_path)
    current_workbook.close()

    # Print Cell F9 Value to Cell H9
    # current_workbook = load_workbook(file_path)
    # current_worksheet = current_workbook.active
    # current_worksheet['H9'] = cell_f9_value
    # current_workbook.save(file_path)
    # current_workbook.close()

    # # Write the values to the output worksheet
    output_worksheet.cell(row=start_row, column=start_column, value=cell_f9_value)

    # Increment the row for the next iteration
    start_row += 1

    # Print the values
    print(f"File: {file}")
    print(f"Loyaty amount: {cell_f9_value}")
    print()

    # Save the changes and close the output workbook
    output_workbook.save(output_file)
    output_workbook.close()

# Add a user input prompt to keep the console window open
print()
print("Finished processing")
input("Press Enter to exit...")